import io
import tempfile
import unittest
from pathlib import Path

import vault.db as vault_db
from vault.export import to_docx, to_xlsx
from vault.extract import extract_text
from vault.ingest import ingest_inbox, ingest_text, ingest_uploaded, safe_upload_name
from vault.menus import ANY, describe_selection, menu_options, search_kwargs
from vault.menus import FIELD_MENUS
from vault.parser import parse_vault_profile
from vault.search import search_profiles


SAMPLE = """CONSENT:YES
Gender: Male
Name: SARVEPALLI VENKATESH BABU
Father name: SREE HARI RAO SARVEPALLI
Mothername: MADHAVI SARVEPALLI
DOB: 30-07-1995
Nakshatram: Maka, 4 PADAM
Rashi: SIMHA
Place of Birth: VIJAYAWADA
Education: B. TECH. (Civil)
ContactNumber: 9391073474, 6281790680, 9491728910
"""


class VaultTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.old = vault_db.DB_PATH
        vault_db.DB_PATH = Path(self.tmp.name) / "vault.sqlite3"
        self.inbox = Path(self.tmp.name) / "inbox"
        self.inbox.mkdir()

    def tearDown(self):
        vault_db.DB_PATH = self.old
        self.tmp.cleanup()

    def test_parse_keeps_phone_and_maka_padam(self):
        p = parse_vault_profile(SAMPLE)
        self.assertIsNotNone(p)
        self.assertEqual(p["gender"], "male")
        self.assertEqual(p["nakshatram"], "magha")
        self.assertEqual(p["padam"], "4")
        self.assertEqual(p["rashi"], "simha")
        self.assertIn("9391073474", p["phone"])

    def test_ingest_and_search_full_details(self):
        yes, skipped = ingest_text(SAMPLE, source_file="sample.txt", source_kind="text")
        self.assertEqual(yes, 1)
        rows = search_profiles(stars="Maka", padam="4", gender="Male", phone="93910")
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["name"], "SARVEPALLI VENKATESH BABU")
        self.assertIn("9391073474", rows[0]["phone"])
        self.assertIn("SREE HARI", rows[0]["father_name"])

    def test_inbox_ingest(self):
        (self.inbox / "one.txt").write_text(SAMPLE, encoding="utf-8")
        stats = ingest_inbox(self.inbox, permission_confirmed=True)
        self.assertEqual(stats["accepted"], 1)
        self.assertEqual(vault_db.count_profiles(), 1)
        with self.assertRaises(ValueError):
            ingest_inbox(self.inbox, permission_confirmed=False)

    def test_upload_from_desktop_bytes(self):
        stats = ingest_uploaded(
            [("../../outside.txt", SAMPLE.encode())],
            permission_confirmed=True,
            dest_dir=self.inbox,
        )
        self.assertEqual(stats["accepted"], 1)
        self.assertTrue((self.inbox / "outside.txt").is_file())
        self.assertFalse((Path(self.tmp.name) / "outside.txt").exists())
        self.assertEqual(vault_db.count_profiles(), 1)
        again = ingest_uploaded(
            [("outside.txt", SAMPLE.encode())],
            permission_confirmed=True,
            dest_dir=self.inbox,
        )
        self.assertEqual(again["accepted"], 1)
        self.assertEqual(vault_db.count_profiles(), 1)
        with self.assertRaises(ValueError):
            ingest_uploaded([("a.txt", b"x")], permission_confirmed=False, dest_dir=self.inbox)

    def test_upload_skips_unsupported_without_writing(self):
        stats = ingest_uploaded(
            [("notes.exe", b"MZ")],
            permission_confirmed=True,
            dest_dir=self.inbox,
        )
        self.assertEqual(stats["unsupported"], 1)
        self.assertEqual(stats["accepted"], 0)
        self.assertFalse((self.inbox / "notes.exe").exists())
        self.assertEqual(safe_upload_name("../../notes.exe"), "notes.exe")

    def test_docx_round_trip(self):
        from docx import Document

        path = self.inbox / "bio.docx"
        doc = Document()
        for line in SAMPLE.splitlines():
            doc.add_paragraph(line)
        doc.save(path)
        text, kind = extract_text(path)
        self.assertEqual(kind, "docx")
        self.assertIn("VENKATESH", text)
        stats = ingest_uploaded(
            [("bio.docx", path.read_bytes())],
            permission_confirmed=True,
            dest_dir=self.inbox / "uploads",
        )
        self.assertEqual(stats["accepted"], 1)

    def test_menus_drive_search_without_typed_query(self):
        ingest_text(SAMPLE, source_file="sample.txt", source_kind="text")
        profiles = vault_db.load_all()
        star_menu = next(spec for spec in FIELD_MENUS if spec["key"] == "stars")
        options = menu_options(profiles, star_menu)
        self.assertEqual(options[0], ANY)
        self.assertIn("Magha", options)
        self.assertIn("Nakshatram: Magha", describe_selection({"stars": "Magha"}, ANY))
        self.assertEqual(describe_selection({}, ANY), "All stored profiles")
        matched = search_profiles(**search_kwargs({"stars": "Magha", "gender": "Male"}, ANY))
        self.assertEqual(len(matched), 1)
        self.assertIn("9391073474", matched[0]["phone"])
        too_young = search_profiles(**search_kwargs({"stars": "Magha"}, "18-22"))
        self.assertEqual(too_young, [])
        in_range = search_profiles(**search_kwargs({"stars": "Magha"}, "28-35"))
        self.assertEqual(len(in_range), 1)

    def test_submit_exports_keep_phone_and_block_formulas(self):
        ingest_text(SAMPLE, source_file="sample.txt", source_kind="text")
        rows = search_profiles(**search_kwargs({"stars": "Magha"}, ANY))
        rows.append({"name": "=1+1", "phone": "9999999999", "nakshatram": "rohini"})
        blob = to_xlsx(rows, "Nakshatram: Magha")
        self.assertTrue(blob.startswith(b"PK"))
        from openpyxl import load_workbook

        book = load_workbook(io.BytesIO(blob))
        sheet = book["Profiles"]
        values = [cell.value for row in sheet.iter_rows() for cell in row]
        self.assertIn("9391073474, 6281790680, 9491728910", values)
        self.assertIn("'=1+1", values)
        word = to_docx(rows[:1], "Nakshatram: Magha")
        from docx import Document

        doc = Document(io.BytesIO(word))
        table_text = [
            " ".join(cell.text for cell in row.cells)
            for table in doc.tables
            for row in table.rows
        ]
        self.assertTrue(any("9391073474" in line for line in table_text))
        self.assertIn("Wedding Vault", "\n".join(p.text for p in doc.paragraphs))

    def test_desktop_screen_has_import_and_submit_menus(self):
        source = (Path(__file__).resolve().parents[1] / "vault" / "app.py").read_text(encoding="utf-8")
        self.assertIn('form_submit_button("Submit"', source)
        self.assertIn("file_uploader(", source)
        self.assertIn("Import files", source)
        self.assertNotIn("text_input(", source)
        self.assertNotIn("st.slider(", source)


if __name__ == "__main__":
    unittest.main()
