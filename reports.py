"""Local XLSX, DOCX, PDF exports. PDF conversion requires LibreOffice installed."""
from __future__ import annotations

import io
import os
import shutil
import subprocess
import tempfile
from datetime import datetime
from pathlib import Path

FIELDS = [
    ("name", "Bride Name"), ("age", "Age"), ("dob", "DOB"),
    ("nakshatram", "Nakshatram"), ("rashi", "Rashi"),
    ("gotram", "Gotram"), ("saka", "Saka"),
    ("birthplace", "Place of Birth"), ("birthtime", "Birth Time"),
    ("education", "Education"), ("location", "Current Location"),
    ("source_name", "Source Group"), ("source_ref", "Message ID"),
    ("status", "Match Status"), ("duplicate_of", "Potential Duplicate"),
]

def to_xlsx(records: list[dict], filter_summary: str = "") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.page import PageMargins

    book = Workbook()
    info = book.active
    info.title = "Overview"
    info["A1"] = "MARRIAGE PROFILE MATCH REPORT"
    info["A1"].font = Font(bold=True, size=17, color="14375A")
    info["A3"] = "Generated"
    info["B3"] = datetime.now().strftime("%d %b %Y %I:%M %p")
    info["A4"] = "Applied filters"
    info["B4"] = filter_summary or "None"
    info["A5"] = "Profiles in report"
    info["B5"] = len(records)
    info["A6"] = "Strict matches"
    info["B6"] = sum(r.get("status") == "Match" for r in records)
    info["A7"] = "Needs review"
    info["B7"] = sum(r.get("status") == "Needs review" for r in records)
    info["A9"] = "Privacy: Share this file only with authorized persons; no contact numbers are collected."
    info.column_dimensions["A"].width = 27
    info.column_dimensions["B"].width = 62
    info["B4"].alignment = Alignment(wrap_text=True)
    info.row_dimensions[4].height = 45

    ws = book.create_sheet("Profiles")
    ws.append([label for _, label in FIELDS])
    header_fill = PatternFill("solid", fgColor="183A58")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True, size=10)
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 32
    for index, record in enumerate(records, 2):
        vals = []
        for key, _ in FIELDS:
            value = record.get(key)
            if key == "duplicate_of":
                value = "Review duplicate" if value else ""
            elif value is None:
                value = "Not provided"
            # Protect spreadsheet software against formula injection from user supplied text.
            if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")):
                value = "'" + value
            vals.append(value)
        ws.append(vals)
        for c in ws[index]:
            c.fill = PatternFill("solid", fgColor="EEF4F8" if index % 2 == 0 else "FFFFFF")
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.font = Font(size=10, color="243749")
        ws.row_dimensions[index].height = 33
    widths = [23, 9, 16, 17, 17, 18, 20, 24, 17, 20, 20, 30, 15, 18, 19]
    for j, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = width
    ws.freeze_panes = "A2"
    last_col = get_column_letter(len(FIELDS))
    ws.auto_filter.ref = f"A1:{last_col}{max(1, len(records) + 1)}"
    ws.sheet_view.showGridLines = False
    info.sheet_view.showGridLines = False
    ws.print_options.horizontalCentered = True
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_margins = PageMargins(left=.25, right=.25, top=.5, bottom=.5, header=.2, footer=.2)
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def to_docx(records: list[dict], filter_summary: str = "") -> bytes:
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(.7)
    sec.bottom_margin = Inches(.7)
    sec.left_margin = Inches(.8)
    sec.right_margin = Inches(.8)
    doc.styles["Normal"].font.name = "Noto Sans"
    doc.styles["Normal"].font.size = Pt(10)
    title = doc.add_heading("Marriage Profile Match Report", 0)
    title.style = "Title"
    doc.add_paragraph(f"Generated: {datetime.now():%d %B %Y}    |    Profiles: {len(records)}")
    doc.add_paragraph(f"Filters: {filter_summary or 'None'}")
    doc.add_paragraph("Confidential — for authorized matchmaking review only. Fields not provided are marked explicitly.")
    for ix, record in enumerate(records, 1):
        name = record.get("name") or "Unnamed profile"
        doc.add_heading(f"{ix}. {name}  [{record.get('status','Unreviewed')}]", level=2)
        table = doc.add_table(rows=0, cols=2)
        table.style = "Light Shading Accent 1"
        for key, label in FIELDS:
            if key in ("name", "duplicate_of", "status"):
                continue
            row = table.add_row().cells
            row[0].text = label
            row[1].text = str(record.get(key) if record.get(key) not in (None, "") else "Not provided")
        if record.get("duplicate_of"):
            doc.add_paragraph("Potential duplicate — verify before using this profile.")
    if not records:
        doc.add_paragraph("No matching profiles. Adjust filters or import consented profiles.")
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def to_pdf(records: list[dict], filter_summary: str = "") -> bytes:
    """DOCX -> PDF via LibreOffice; respects Indic shaping better than default PDF fonts."""
    office = shutil.which("libreoffice") or shutil.which("soffice")
    if not office:
        raise RuntimeError("PDF export requires LibreOffice: install it, ensure soffice/libreoffice is in PATH, and install Telugu fonts such as Noto Sans Telugu.")
    with tempfile.TemporaryDirectory(prefix="matchreport_") as directory:
        path = Path(directory)
        docx_path = path / "match_report.docx"
        docx_path.write_bytes(to_docx(records, filter_summary))
        result = subprocess.run([office, "-env:UserInstallation=file://" + str(path / 'lo_profile'),
              "--headless", "--convert-to", "pdf", "--outdir", str(path), str(docx_path)],
              text=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=120)
        target = path / "match_report.pdf"
        if result.returncode != 0 or not target.exists():
            raise RuntimeError(f"PDF conversion failed: {result.stderr or result.stdout}")
        return target.read_bytes()
