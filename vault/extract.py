"""Extract plain text from dumped files for Vault ingest."""
from __future__ import annotations

import json
from pathlib import Path

TEXT_SUFFIXES = {".txt", ".text", ".md", ".csv"}
EXCEL_SUFFIXES = {".xlsx", ".xlsm"}
PDF_SUFFIXES = {".pdf"}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff", ".bmp"}
JSON_SUFFIXES = {".json"}
DOCX_SUFFIXES = {".docx"}
SUPPORTED_SUFFIXES = (
    TEXT_SUFFIXES | EXCEL_SUFFIXES | PDF_SUFFIXES | IMAGE_SUFFIXES | JSON_SUFFIXES | DOCX_SUFFIXES
)


def extract_text(path: Path) -> tuple[str, str]:
    """Return (text, source_kind). Raises ValueError if unsupported/unreadable."""
    suffix = path.suffix.casefold()
    if suffix in TEXT_SUFFIXES or suffix == "":
        return path.read_text(encoding="utf-8-sig", errors="replace"), "text"
    if suffix in JSON_SUFFIXES:
        return _from_json(path), "json"
    if suffix in EXCEL_SUFFIXES:
        return _from_excel(path), "excel"
    if suffix in PDF_SUFFIXES:
        return _from_pdf(path), "pdf"
    if suffix in IMAGE_SUFFIXES:
        return _from_image(path), "image"
    if suffix in DOCX_SUFFIXES:
        return _from_docx(path), "docx"
    raise ValueError(f"Unsupported file type: {suffix or '(none)'}")


def _from_json(path: Path) -> str:
    obj = json.loads(path.read_text(encoding="utf-8-sig", errors="replace"))
    if isinstance(obj, dict) and isinstance(obj.get("messages"), list):
        chunks = []
        for message in obj["messages"]:
            if not isinstance(message, dict):
                continue
            parts = message.get("text", "")
            if isinstance(parts, list):
                raw = "".join(p if isinstance(p, str) else str(p.get("text", "")) for p in parts)
            else:
                raw = str(parts)
            if raw.strip():
                chunks.append(raw.strip())
        return "\n\n".join(chunks)
    return path.read_text(encoding="utf-8-sig", errors="replace")


def _from_excel(path: Path) -> str:
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=True)
    lines: list[str] = []
    for sheet in book.worksheets:
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c not in (None, "")]
            if not cells:
                continue
            if len(cells) == 2:
                lines.append(f"{cells[0]}: {cells[1]}")
            else:
                lines.append(" | ".join(cells))
        lines.append("")
    return "\n".join(lines)


def _from_pdf(path: Path) -> str:
    try:
        import fitz  # PyMuPDF
    except ImportError as exc:
        raise ValueError("PDF support needs pymupdf. pip install pymupdf") from exc
    doc = fitz.open(path)
    try:
        return "\n".join(page.get_text("text") for page in doc)
    finally:
        doc.close()


def _from_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise ValueError("Word support needs python-docx. pip install python-docx") from exc
    doc = Document(str(path))
    lines = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()]
            if not cells:
                continue
            if len(cells) == 2:
                lines.append(f"{cells[0]}: {cells[1]}")
            else:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def _from_image(path: Path) -> str:
    try:
        from PIL import Image
        import pytesseract
    except ImportError as exc:
        raise ValueError(
            "Image OCR needs pillow + pytesseract (and system tesseract binary). "
            "pip install pillow pytesseract"
        ) from exc
    image = Image.open(path)
    text = pytesseract.image_to_string(image)
    if not text.strip():
        raise ValueError("OCR returned no text from image")
    return text
